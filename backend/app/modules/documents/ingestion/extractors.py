"""Pipeline step 2: file bytes -> plain text sections.

Validation happens BEFORE anything is stored: never trust the filename or the size the client says.
"""

import io
from dataclasses import dataclass
from pathlib import PurePath
from pypdf import PdfReader
from docx import Document
import csv

ALLOWED_TYPES = ("pdf", "txt", "md", "csv", "docx")
MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB


class UnsupportedFileError(ValueError):
    """Wrong extension or too big. The API turns it into a 4xx."""


class EmptyDocumentError(ValueError):
    """The file is valid but contains no extractable text (e.g. a scanned PDF)."""


@dataclass
class Section:
    text: str
    page: int | None = None  # 1-based page number when the format has pages (PDF)


def safe_filename(filename: str) -> str:
    """Only the last path component: 'C:\\x\\..\\a.pdf' or '../../a.pdf' becomes 'a.pdf'."""
    return PurePath(filename.replace("\\", "/")).name


def validate_upload(filename: str, size_bytes: int) -> str:
    """Return the file type ('pdf', 'txt', ...) or raise UnsupportedFileError."""
    name = safe_filename(filename or "")
    extension = name.rsplit(".", 1)[-1].lower() if "." in name else ""

    if extension not in ALLOWED_TYPES:
        raise UnsupportedFileError(
            f"Unsupported file type. Allowed: {', '.join(ALLOWED_TYPES)}"
        )

    if size_bytes == 0:
        raise UnsupportedFileError("The file is empty")

    if size_bytes > MAX_FILE_BYTES:
        raise UnsupportedFileError(
            f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB"
        )

    return extension


def extract(file_type: str, data: bytes) -> list[Section]:
    extractor = _EXTRACTORS[file_type]
    sections = [s for s in extractor(data) if s.text.strip()]

    if not sections:
        raise EmptyDocumentError("No text could be extracted from the file")

    return sections


# --- one function per format -------------------------------------------------


def _from_text(data: bytes) -> list[Section]:
    """TXT and Markdown are already text: just decode them."""
    try:
        return [Section(data.decode("utf-8"))]
    except UnicodeDecodeError:
        return [Section(data.decode("latin-1"))]  # Windows files are often not UTF-8


def _from_pdf(data: bytes) -> list[Section]:

    reader = PdfReader(io.BytesIO(data))
    sections = []

    for page_number, page in enumerate(reader.pages, start=1):
        sections.append(
            Section(
                text=page.extract_text() or "",
                page=page_number,
            )
        )

    return sections


def _from_docx(data: bytes) -> list[Section]:
    document = Document(io.BytesIO(data))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]

            if any(cells):
                parts.append(" | ".join(cells))

    return [Section("\n".join(parts))]


def _from_csv(data: bytes) -> list[Section]:

    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")

    reader = csv.DictReader(io.StringIO(text))

    lines = []

    for row in reader:

        fields = [
            f"{column}: {value if value is not None else ''}"
            for column, value in row.items()
            if column is not None
        ]

        lines.append("; ".join(fields))

    return [Section("\n".join(lines))]


_EXTRACTORS = {
    "txt": _from_text,
    "md": _from_text,
    "pdf": _from_pdf,
    "docx": _from_docx,
    "csv": _from_csv,
}
