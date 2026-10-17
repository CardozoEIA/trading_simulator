# tests/modules/documents/test_extractors.py
import pytest

import io

from docx import Document
from pypdf import PdfWriter
from pathlib import Path

from app.modules.documents.ingestion.extractors import (
    MAX_FILE_BYTES,
    EmptyDocumentError,
    UnsupportedFileError,
    extract,
    safe_filename,
    validate_upload,
)


def test_accepts_allowed_extensions_case_insensitive():
    assert validate_upload("Notes.PDF", 100) == "pdf"
    assert validate_upload("glossary.md", 100) == "md"


@pytest.mark.parametrize("name", ["virus.exe", "noextension", "archive.zip", ""])
def test_rejects_other_types(name):
    with pytest.raises(UnsupportedFileError):
        validate_upload(name, 100)


def test_rejects_empty_and_oversized_files():
    with pytest.raises(UnsupportedFileError):
        validate_upload("a.txt", 0)

    with pytest.raises(UnsupportedFileError):
        validate_upload("a.txt", MAX_FILE_BYTES + 1)


def test_safe_filename_drops_directories():
    assert safe_filename("../../etc/passwd.txt") == "passwd.txt"
    assert safe_filename("C:\\Users\\me\\a.pdf") == "a.pdf"


def test_txt_utf8_and_latin1():
    assert extract("txt", "Rentabilidad ñ".encode("utf-8"))[0].text == "Rentabilidad ñ"
    assert (
        extract("txt", "Rentabilidad ñ".encode("latin-1"))[0].text == "Rentabilidad ñ"
    )


def test_blank_text_is_rejected():
    with pytest.raises(EmptyDocumentError):
        extract("txt", b"   \n  ")


def test_csv_rows_keep_their_column_names():
    data = b"date,open,close\n" b"2024-01-02,4700,4750\n" b"2024-01-03,4750,4800\n"

    sections = extract("csv", data)

    assert len(sections) == 1

    text = sections[0].text

    assert "date: 2024-01-02" in text
    assert "close: 4750" in text
    assert "date: 2024-01-03" in text
    assert "close: 4800" in text


def test_docx_paragraphs_are_extracted():
    doc = Document()

    doc.add_paragraph("Sharpe ratio")
    doc.add_paragraph("Risk adjusted return")

    buffer = io.BytesIO()
    doc.save(buffer)

    sections = extract("docx", buffer.getvalue())

    assert len(sections) == 1
    assert "Sharpe ratio" in sections[0].text
    assert "Risk adjusted return" in sections[0].text


def test_pdf_sections_carry_page_numbers():
    pdf_path = Path(__file__).parent.parent.parent / "fixtures" / "two_pages.pdf"

    data = pdf_path.read_bytes()

    sections = extract("pdf", data)

    assert len(sections) == 2

    assert sections[0].page == 1
    assert sections[1].page == 2

    assert "Sharpe" in sections[0].text
    assert "drawdown" in sections[1].text
